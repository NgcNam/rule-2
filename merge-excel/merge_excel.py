import sys
import glob
import os
import openpyxl

def main():
    if len(sys.argv) < 2:
        print("Cach dung: python merge_excel.py <ten_folder>")
        sys.exit(1)

    folder = sys.argv[1]
    if not os.path.isdir(folder):
        print(f"Khong tim thay folder: {folder}")
        sys.exit(1)

    # Ten file ket qua = ten folder
    out_name = os.path.basename(os.path.normpath(folder)) + ".xlsx"

    files = glob.glob(os.path.join(folder, "*.xlsx"))
    # Bo qua file ket qua neu chay lai
    files = [f for f in files if os.path.basename(f) != out_name]

    if not files:
        print("Khong co file .xlsx nao trong folder.")
        sys.exit(1)

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.append(["nội dung", "value"])

    total = 0
    for f in files:
        name = os.path.basename(f).lower()
        # Kiem tra "nonspam" truoc vi no cung chua "spam"
        if "nonspam" in name:
            value = "nonspam"
        elif "spam" in name:
            value = "spam"
        else:
            print(f"Bo qua (ten khong chua spam/nonspam): {f}")
            continue

        wb = openpyxl.load_workbook(f, data_only=True)
        sheet_name = next((s for s in wb.sheetnames if s.lower() == "social"), None)
        if sheet_name is None:
            print(f"Bo qua (khong co sheet 'social'): {f}")
            continue
        ws = wb[sheet_name]

        count = 0
        # Cot D = cot 4, bat dau tu dong 12
        for row in range(12, ws.max_row + 1):
            content = ws.cell(row=row, column=4).value
            if content is None or str(content).strip() == "":
                continue
            out_ws.append([content, value])
            count += 1

        total += count
        print(f"{os.path.basename(f)} -> {value}: {count} dong")

    out_path = os.path.join(folder, out_name)
    out_wb.save(out_path)
    print(f"\nDa luu: {out_path} (tong {total} dong)")

if __name__ == "__main__":
    main()
