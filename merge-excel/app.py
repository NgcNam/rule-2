import io
import openpyxl
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse

app = FastAPI(title="Merge Excel")


# Cau hinh cho tung loai file:
#   sheet   : ten sheet (so khop khong phan biet hoa/thuong)
#   min_row : dong bat dau lay du lieu
#   col     : so thu tu cot (D=4, G=7)
FILE_TYPES = {
    "social": {"sheet": "social", "min_row": 12, "col": 4},          # D12 tro xuong, sheet Social
    "admin_database": {"sheet": "sheet1", "min_row": 7, "col": 7},    # G7 tro xuong, sheet Sheet1
}


def extract_rows(file_bytes: bytes, filename: str, value: str, file_type: str):
    """Lay du lieu theo cau hinh loai file, gan nhan theo value truyen vao."""
    cfg = FILE_TYPES.get(file_type)
    if cfg is None:
        raise HTTPException(status_code=400, detail=f"Loai file khong hop le: '{file_type}'.")

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet_name = next((s for s in wb.sheetnames if s.lower() == cfg["sheet"]), None)
    if sheet_name is None:
        raise HTTPException(
            status_code=400,
            detail=f"File '{filename}' khong co sheet '{cfg['sheet']}'.",
        )
    ws = wb[sheet_name]

    rows = []
    for row in ws.iter_rows(min_row=cfg["min_row"], min_col=cfg["col"], max_col=cfg["col"], values_only=True):
        content = row[0]
        if content is None or str(content).strip() == "":
            continue
        rows.append((content, value))
    wb.close()
    return rows


INDEX_HTML = """<!doctype html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Merge Excel</title>
<style>
  * { box-sizing: border-box; }
  body { font-family: system-ui, -apple-system, "Segoe UI", sans-serif; background:#f4f5f7; margin:0; padding:40px 16px; color:#1a1a1a; }
  .card { max-width:560px; margin:0 auto; background:#fff; border-radius:14px; padding:32px; box-shadow:0 4px 24px rgba(0,0,0,.08); }
  h1 { font-size:22px; margin:0 0 6px; }
  p.sub { margin:0 0 24px; color:#666; font-size:14px; }
  label { display:block; font-weight:600; font-size:14px; margin:18px 0 6px; }
  input[type=file] { width:100%; padding:12px; border:1.5px dashed #cbd2dc; border-radius:10px; background:#fafbfc; font-size:14px; }
  select { width:100%; padding:11px; border:1.5px solid #cbd2dc; border-radius:10px; background:#fff; font-size:14px; margin-top:6px; }
  button { margin-top:26px; width:100%; padding:14px; border:0; border-radius:10px; background:#2563eb; color:#fff; font-size:16px; font-weight:600; cursor:pointer; }
  button:hover { background:#1d4ed8; }
  button:disabled { background:#9db4e8; cursor:not-allowed; }
  .note { margin-top:18px; font-size:12.5px; color:#888; line-height:1.6; }
  #status { margin-top:18px; font-size:14px; min-height:20px; }
  .ok { color:#15803d; } .err { color:#dc2626; }
</style>
</head>
<body>
  <div class="card">
    <h1>Merge Excel</h1>
    <p class="sub">Chọn file cho từng ô. File ở ô <b>spam</b> được gán value = spam, ô <b>nonspam</b> được gán value = nonspam.</p>
    <form id="form">
      <label>File spam</label>
      <input type="file" name="file_spam" accept=".xlsx" required>
      <select name="type_spam">
        <option value="social">Social (cột D, sheet Social, từ dòng 12)</option>
        <option value="admin_database">Admin database (cột G, Sheet1, từ dòng 7)</option>
      </select>
      <label>File nonspam</label>
      <input type="file" name="file_nonspam" accept=".xlsx" required>
      <select name="type_nonspam">
        <option value="social">Social (cột D, sheet Social, từ dòng 12)</option>
        <option value="admin_database">Admin database (cột G, Sheet1, từ dòng 7)</option>
      </select>
      <label>Tên file kết quả</label>
      <input type="text" name="outname" value="merged" style="width:100%;padding:12px;border:1.5px solid #cbd2dc;border-radius:10px;font-size:14px;">
      <button type="submit" id="btn">Merge & Tải về</button>
    </form>
    <div id="status"></div>
    <div class="note">
      Mỗi file chọn loại tương ứng. <b>Social</b>: cột D sheet Social từ dòng 12.
      <b>Admin database</b>: cột G Sheet1 từ dòng 7. Value gán theo ô (spam / nonspam).
      Cột kết quả: <b>nội dung</b>, <b>value</b>.
    </div>
  </div>
<script>
const form = document.getElementById('form');
const btn = document.getElementById('btn');
const status = document.getElementById('status');
form.addEventListener('submit', async (e) => {
  e.preventDefault();
  status.textContent = ''; status.className = '';
  btn.disabled = true; btn.textContent = 'Đang xử lý...';
  try {
    const data = new FormData(form);
    const res = await fetch('/merge', { method: 'POST', body: data });
    if (!res.ok) {
      const j = await res.json().catch(() => ({detail: 'Lỗi không xác định'}));
      throw new Error(j.detail || ('HTTP ' + res.status));
    }
    const blob = await res.blob();
    const dispo = res.headers.get('Content-Disposition') || '';
    const m = dispo.match(/filename="?([^"]+)"?/);
    const fname = m ? m[1] : 'merged.xlsx';
    const url = URL.createObjectURL(blob);
    const a = document.createElement('a');
    a.href = url; a.download = fname; a.click();
    URL.revokeObjectURL(url);
    status.textContent = '✔ Đã tạo file: ' + fname;
    status.className = 'ok';
  } catch (err) {
    status.textContent = '✖ ' + err.message;
    status.className = 'err';
  } finally {
    btn.disabled = false; btn.textContent = 'Merge & Tải về';
  }
});
</script>
</body>
</html>
"""


@app.get("/", response_class=HTMLResponse)
def index():
    return INDEX_HTML


@app.post("/merge")
async def merge(
    file_spam: UploadFile = File(...),
    file_nonspam: UploadFile = File(...),
    type_spam: str = Form("social"),
    type_nonspam: str = Form("social"),
    outname: str = Form("merged"),
):
    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.append(["nội dung", "value"])

    for f, value, ftype in ((file_spam, "spam", type_spam), (file_nonspam, "nonspam", type_nonspam)):
        data = await f.read()
        rows = extract_rows(data, f.filename, value, ftype)
        for content, val in rows:
            out_ws.append([content, val])

    buf = io.BytesIO()
    out_wb.save(buf)
    buf.seek(0)

    safe = (outname or "merged").strip() or "merged"
    if not safe.lower().endswith(".xlsx"):
        safe += ".xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe}"'},
    )


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
