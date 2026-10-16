# -*- coding: utf-8 -*-
"""
STEP 6 — Backtest toàn bộ + xuất Excel theo khung format.xlsx.

Áp rulebase (step5) lên train & test, tính:
  - Chỉ số hệ thống trên TEST (precision/recall/f1/accuracy) + kiểm tra ngưỡng.
  - TP/FP/TN/FN từng rule cho CẢ train lẫn test.
Xuất workbook gồm 6 sheet: Summary, RuleStats, AllResults, FalsePositives,
FalseNegatives, RuleExamples. (Trong RuleStats, các keyword ngăn cách bằng dấu ",")

INPUT : step1_segmented.json, step5_final_rules.json
OUTPUT: report.xlsx

Chạy:
  python step6_export.py --step1 ../out/step1_segmented.json \
      --step5 ../out/step5_final_rules.json --output ../out/report.xlsx
"""
import argparse

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import (DEFAULTS, classify, kw_display, match_tokens, prf,
                    read_json, rule_matches)

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)
SECTION_FONT = Font(bold=True, size=11, color="1F4E78")


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def rule_counts(rule, docs, spam_lbl):
    """(tp, fp, tn, fn) của 1 rule trên tập docs."""
    tp = fp = tn = fn = 0
    for d in docs:
        m = rule_matches(rule, d["_mtok"])
        is_spam = d["label"] == spam_lbl
        if m and is_spam:
            tp += 1
        elif m and not is_spam:
            fp += 1
        elif not m and is_spam:
            fn += 1
        else:
            tn += 1
    return tp, fp, tn, fn


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--step1", required=True)
    ap.add_argument("--step5", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    s1 = read_json(args.step1)
    s5 = read_json(args.step5)
    spam_lbl = DEFAULTS["spam_label"]

    docs = s1["documents"]
    for d in docs:                      # token match nhất quán với step2/3
        d["_mtok"] = match_tokens(d["segmented_text"])
    train = [d for d in docs if d["split"] == "train"]
    test = [d for d in docs if d["split"] == "test"]
    spam_rules = s5["spam_rules"]
    nonspam_rules = s5["nonspam_rules"]

    # ---- phân loại toàn bộ document ----
    results = []
    conf_conflict = conf_nomatch = 0
    for d in docs:
        pred, detail, ms, mn = classify(spam_rules, nonspam_rules, d["_mtok"])
        if d["split"] == "test":
            if detail == "uncertain_spam":
                conf_conflict += 1
            elif detail == "uncertain":
                conf_nomatch += 1
        results.append({
            "doc": d, "pred": pred, "detail": detail,
            "spam_ids": ms, "non_ids": mn,
            "ok": pred == d["label"],
        })

    # ---- chỉ số hệ thống trên TEST ----
    def metrics(subset):
        tp = fp = tn = fn = 0
        for r in subset:
            actual_spam = r["doc"]["label"] == spam_lbl
            pred_spam = r["pred"] == spam_lbl
            if pred_spam and actual_spam:
                tp += 1
            elif pred_spam and not actual_spam:
                fp += 1
            elif not pred_spam and actual_spam:
                fn += 1
            else:
                tn += 1
        return tp, fp, tn, fn

    test_res = [r for r in results if r["doc"]["split"] == "test"]
    train_res = [r for r in results if r["doc"]["split"] == "train"]
    tp, fp, tn, fn = metrics(test_res)
    p, r_, f1 = prf(tp, fp, fn)
    acc = (tp + tn) / max(len(test_res), 1)

    # ============================ BUILD WORKBOOK ============================
    wb = openpyxl.Workbook()

    # ----- Summary -----
    ws = wb.active
    ws.title = "Summary"
    dist = s1["meta"]
    rows = [
        ("SPAM RULE MINING REPORT — Long Châu Dataset", None),
        ("📊 DATASET OVERVIEW", None),
        ("Total", dist["n_total"]),
        ("  — Spam", dist["dist_total"].get(spam_lbl, 0)),
        ("  — NonSpam", dist["dist_total"].get(DEFAULTS["nonspam_label"], 0)),
        ("Dropped (empty/label lỗi)", dist["n_dropped_empty_or_bad_label"]),
        (f"Train ({int(dist['train_ratio']*100)}%)", dist["n_train"]),
        (f"Test ({int((1-dist['train_ratio'])*100)}%)", dist["n_test"]),
        (None, None),
        ("📈 BACKTEST RESULTS (Test Set)", None),
        ("True Positives  (Spam caught)", tp),
        ("False Positives (NonSpam blocked)", fp),
        ("True Negatives  (NonSpam correct)", tn),
        ("False Negatives (Spam missed)", fn),
        (None, None),
        ("🎯 PERFORMANCE METRICS", None),
        ("Precision", f"{p*100:.2f}%  {'✅ PASS' if p>0.99 else '❌ FAIL'}(>99%)"),
        ("Recall", f"{r_*100:.2f}%  {'✅ PASS' if r_>0.30 else '❌ FAIL'}(>30%)"),
        ("F1-Score", f"{f1*100:.2f}%"),
        ("Accuracy", f"{acc*100:.2f}%"),
        (None, None),
        ("📋 RULE INVENTORY", None),
        ("Spam Rules", len(spam_rules)),
        ("NonSpam Rules", len(nonspam_rules)),
        ("Uncertain — conflict (test)", conf_conflict),
        ("Uncertain — no match (test)", conf_nomatch),
        (None, None),
        ("⚙️ CLASSIFICATION LOGIC", None),
        ("Priority 1", "Spam rule matches → SPAM (definitive)"),
        ("Priority 2", "NonSpam only → confirmed NONSPAM"),
        ("Conflict", "Both match → uncertain_spam (treated as SPAM)"),
        ("No match", "→ uncertain (manual review)"),
        (None, None),
        ("🔧 SEGMENTATION", None),
        ("Backend", dist["segmentation_backend"]),
        ("PhoBERT tokenizer", "on" if dist["phobert_tokenizer"] else "off"),
    ]
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(row=i, column=1, value=a)
        if b is not None:
            ws.cell(row=i, column=2, value=b)
    ws.cell(row=1, column=1).font = TITLE_FONT
    for i, (a, b) in enumerate(rows, 1):
        if a and b is None and a[0] in "📊📈🎯📋⚙️🔧":
            ws.cell(row=i, column=1).font = SECTION_FONT
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 40

    # ----- RuleStats -----
    ws = wb.create_sheet("RuleStats")
    headers = ["Rule", "Category", "Description", "OR Keywords", "AND Keywords",
               "NOT Keywords", "#KW",
               "TP(train)", "FP(train)", "TN(train)", "FN(train)",
               "TP(test)", "FP(test)", "TN(test)", "FN(test)",
               "Precision(test)", "Recall(test)", "F1(test)", "Prec✅", "Recall✅"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for rule in spam_rules:
        tr = rule_counts(rule, train, spam_lbl)
        te = rule_counts(rule, test, spam_lbl)
        pp, rr, ff = prf(te[0], te[1], te[3])
        ws.append([
            rule["id"], rule.get("category", ""), rule.get("description", ""),
            ", ".join(kw_display(k) for k in rule.get("or", [])) or None,
            ", ".join(kw_display(k) for k in rule.get("and", [])) or None,
            ", ".join(kw_display(k) for k in rule.get("not", [])) or None,
            rule.get("n_kw", len(rule.get("or", [])) + len(rule.get("and", [])) + len(rule.get("not", []))),
            tr[0], tr[1], tr[2], tr[3],
            te[0], te[1], te[2], te[3],
            round(pp, 4), round(rr, 4), round(ff, 4),
            "✅" if pp > 0.99 else "❌", "✅" if rr > 0.30 else "⚠️",
        ])
    # thêm rule non-spam để tham chiếu (không tính precision spam)
    for rule in nonspam_rules:
        tr = rule_counts(rule, train, spam_lbl)
        te = rule_counts(rule, test, spam_lbl)
        ws.append([
            rule["id"], rule.get("category", ""), rule.get("description", ""),
            ", ".join(kw_display(k) for k in rule.get("or", [])) or None, None, None,
            rule.get("n_kw", len(rule.get("or", []))),
            tr[0], tr[1], tr[2], tr[3], te[0], te[1], te[2], te[3],
            None, None, None, "—", "—",
        ])
    widths = [12, 26, 34, 50, 24, 24, 6, 9, 9, 9, 9, 9, 9, 9, 9, 13, 12, 10, 7, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # ----- AllResults (train + test) -----
    ws = wb.create_sheet("AllResults")
    hdr = ["#", "Split", "Content (300 chars)", "Actual", "Predicted",
           "Detail", "OK?", "Spam Rules", "NonSpam Rules"]
    ws.append(hdr)
    style_header(ws, 1, len(hdr))
    for i, r in enumerate(results, 1):
        d = r["doc"]
        ws.append([
            i, d["split"], d["content"][:300], d["label"], r["pred"],
            r["detail"], "✅" if r["ok"] else "❌",
            ", ".join(r["spam_ids"]) or "-", ", ".join(r["non_ids"]) or "-",
        ])
    for col, w in zip("ABCDEFGHI", [6, 7, 70, 9, 10, 16, 6, 20, 20]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    # ----- FalsePositives (test set: NonSpam bị chặn nhầm) -----
    ws = wb.create_sheet("FalsePositives")
    fps = [r for r in test_res if r["pred"] == spam_lbl and r["doc"]["label"] != spam_lbl]
    ws.append([f"FALSE POSITIVES — NonSpam bị chặn nhầm (test): {len(fps)} cases"])
    ws.append(["#", "Content", "Actual", "Predicted", "Matched Spam Rule(s)"])
    style_header(ws, 2, 5)
    for i, r in enumerate(fps, 1):
        ws.append([i, r["doc"]["content"][:600], r["doc"]["label"],
                   r["pred"], ", ".join(r["spam_ids"])])
    for col, w in zip("ABCDE", [6, 80, 9, 10, 24]):
        ws.column_dimensions[col].width = w

    # ----- FalseNegatives (test set: Spam bị bỏ sót) -----
    ws = wb.create_sheet("FalseNegatives")
    fns = [r for r in test_res if r["pred"] != spam_lbl and r["doc"]["label"] == spam_lbl]
    ws.append([f"FALSE NEGATIVES — Spam bị bỏ sót (test): {len(fns)} cases"])
    ws.append(["#", "Content", "Actual", "Predicted", "Detail"])
    style_header(ws, 2, 5)
    for i, r in enumerate(fns, 1):
        ws.append([i, r["doc"]["content"][:600], r["doc"]["label"], r["pred"], r["detail"]])
    for col, w in zip("ABCDE", [6, 80, 9, 10, 16]):
        ws.column_dimensions[col].width = w

    # ----- RuleExamples -----
    ws = wb.create_sheet("RuleExamples")
    ws.append(["Rule", "Match Type", "Content (400 chars)"])
    style_header(ws, 1, 3)
    by_rule_tp = {r["id"]: [] for r in spam_rules}
    by_rule_fp = {r["id"]: [] for r in spam_rules}
    for r in results:
        for rid in r["spam_ids"]:
            if r["doc"]["label"] == spam_lbl and len(by_rule_tp[rid]) < 3:
                by_rule_tp[rid].append(r["doc"]["content"][:400])
            elif r["doc"]["label"] != spam_lbl and len(by_rule_fp[rid]) < 2:
                by_rule_fp[rid].append(r["doc"]["content"][:400])
    for rule in spam_rules:
        rid = rule["id"]
        for ex in by_rule_tp[rid]:
            ws.append([rid, "✅ TRUE POSITIVE", ex])
        for ex in by_rule_fp[rid]:
            ws.append([rid, "❌ FALSE POSITIVE", ex])
    for col, w in zip("ABC", [12, 20, 90]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    wb.save(args.output)
    print(f"[step6] OK -> {args.output}")
    print(f"[step6] TEST: TP={tp} FP={fp} TN={tn} FN={fn} | "
          f"Precision={p*100:.2f}% Recall={r_*100:.2f}% F1={f1*100:.2f}% Acc={acc*100:.2f}%")
    print(f"[step6] Precision>99%: {'PASS' if p>0.99 else 'FAIL'} | "
          f"Recall>30%: {'PASS' if r_>0.30 else 'FAIL'}")


if __name__ == "__main__":
    main()
