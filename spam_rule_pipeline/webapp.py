# -*- coding: utf-8 -*-
"""
webapp.py — Web local (Flask) bọc quanh pipeline sinh rulebase spam.

Người dùng: mở web -> UPLOAD file Excel (2 sheet 'spam' & 'nonspam', nhãn = tên
sheet, dữ liệu từ dòng 2) -> hệ thống chạy 6 bước -> XEM kết quả (precision/recall
+ toàn bộ rulebase OR/AND/NOT) -> TẢI file report Excel về.

Chạy:
  cd spam_rule_pipeline
  python webapp.py                 # mở http://127.0.0.1:5000
  python webapp.py --port 8000 --host 0.0.0.0

Output mỗi lần chạy nằm trong 1 folder tên = tên file input (tự thêm _v1/_v2 nếu
trùng), đặt tại thư mục gốc dự án (…/rule-2). File bên trong mang đuôi = tên folder.
"""
import argparse
import os
import re
import subprocess
import sys
import threading
import uuid

from flask import (Flask, abort, jsonify, redirect, render_template_string,
                   request, send_file, url_for)

from common import make_output_dir, read_json

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(HERE, ".."))
# Nơi lưu các folder kết quả. Mặc định = thư mục gốc dự án; trong Docker đặt
# RULEBASE_OUTPUT_DIR=/data (mount volume) để kết quả không nằm lẫn với mã nguồn.
OUTPUT_ROOT = os.environ.get("RULEBASE_OUTPUT_DIR") or PROJECT_ROOT
os.makedirs(OUTPUT_ROOT, exist_ok=True)
UPLOAD_TMP = os.path.join(HERE, "_uploads")
os.makedirs(UPLOAD_TMP, exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024   # 200MB

JOBS = {}                      # job_id -> dict trạng thái
JOBS_LOCK = threading.Lock()

STEP_TITLES = [
    "B1 · Tách từ (PhoBERT/underthesea) + chia train/test",
    "B2 · TF-IDF & spam_score",
    "B3 · Sinh keyword ứng viên (lọc vô nghĩa, khử chứa nhau)",
    "B4 · Gom nhóm ngữ nghĩa + ghép rule + AND/NOT",
    "B5 · Chốt ràng buộc (≤rule/≤kw, khử chứa nhau)",
    "B6 · Backtest + xuất Excel",
]


# ---------------------------------------------------------------- merge ---
def _read_contents(path):
    """Đọc nội dung từ 1 file Excel (1 cột / nhiều cột đều được).

    - Bỏ qua DÒNG ĐẦU (tiêu đề).
    - Chọn cột 'content'/'nội dung' nếu có, ngược lại lấy cột đầu tiên.
    - Trả về list[str] (đã bỏ ô rỗng).
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        ci = 0
        if header is not None:
            hdr = [str(h).strip().lower() if h is not None else "" for h in header]
            for name in ("content", "nội dung", "noi dung"):
                if name in hdr:
                    ci = hdr.index(name)
                    break
        for r in rows:
            val = r[ci] if ci < len(r) else None
            if val is None or str(val).strip() == "":
                continue
            out.append(str(val))
        break          # chỉ đọc sheet đầu tiên của mỗi file
    wb.close()
    return out


def _count_sheets(path):
    """Đếm số dòng dữ liệu (bỏ tiêu đề) cho nhãn spam/nonspam theo TÊN SHEET.

    Sheet tên đúng 'spam' (không phân biệt hoa/thường) -> spam; mọi sheet còn
    lại -> nonspam. Trùng logic step1_segment.py. Trả về (n_spam, n_nonspam).
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    n_spam = n_non = 0
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        ci = 0
        if header is not None:
            hdr = [str(h).strip().lower() if h is not None else "" for h in header]
            for name in ("content", "nội dung", "noi dung"):
                if name in hdr:
                    ci = hdr.index(name)
                    break
        cnt = 0
        for r in rows:
            val = r[ci] if ci < len(r) else None
            if val is not None and str(val).strip() != "":
                cnt += 1
        if str(ws.title).strip().lower() == "spam":
            n_spam += cnt
        else:
            n_non += cnt
    wb.close()
    return n_spam, n_non


def build_combined_excel(spam_path, nonspam_path, out_path):
    """Gộp 2 file Excel (spam + nonspam) thành 1 file 2 sheet 'spam'/'nonspam'
    đúng định dạng pipeline yêu cầu (cột 'content', dữ liệu từ dòng 2)."""
    import openpyxl
    spam_rows = _read_contents(spam_path)
    non_rows = _read_contents(nonspam_path)
    if not spam_rows:
        raise RuntimeError("File SPAM không có dòng dữ liệu nào (bỏ dòng tiêu đề).")
    if not non_rows:
        raise RuntimeError("File NON-SPAM không có dòng dữ liệu nào (bỏ dòng tiêu đề).")
    wb = openpyxl.Workbook()
    ws_spam = wb.active
    ws_spam.title = "spam"
    ws_spam.append(["content"])
    for c in spam_rows:
        ws_spam.append([c])
    ws_non = wb.create_sheet("nonspam")
    ws_non.append(["content"])
    for c in non_rows:
        ws_non.append([c])
    wb.save(out_path)
    return len(spam_rows), len(non_rows)


# --------------------------------------------------------------- pipeline ---
def _set(job_id, **kw):
    with JOBS_LOCK:
        JOBS[job_id].update(kw)


def _step_state(job_id, idx, state, note=""):
    with JOBS_LOCK:
        steps = JOBS[job_id]["steps"]
        steps[idx]["state"] = state
        if note:
            steps[idx]["note"] = note


def run_cmd(job_id, idx, cmd):
    """Chạy 1 step (subprocess), stream stderr để cập nhật tiến độ; trả stdout."""
    _step_state(job_id, idx, "running")
    proc = subprocess.Popen(
        [sys.executable] + cmd, cwd=HERE,
        stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        encoding="utf-8", errors="replace", bufsize=1,
    )
    out_lines = []

    def pump_err():
        for line in proc.stderr:
            line = line.rstrip()
            m = re.search(r"segmented (\d+)/(\d+)", line)
            if m:
                _step_state(job_id, idx, "running",
                            f"tách từ {int(m.group(1)):,}/{int(m.group(2)):,}")
    t = threading.Thread(target=pump_err, daemon=True)
    t.start()
    for line in proc.stdout:
        out_lines.append(line.rstrip())
    proc.wait()
    t.join(timeout=1)
    if proc.returncode != 0:
        raise RuntimeError(f"{STEP_TITLES[idx]} lỗi (exit {proc.returncode})\n"
                           + "\n".join(out_lines[-5:]))
    _step_state(job_id, idx, "done")
    return "\n".join(out_lines)


def worker(job_id, input_path, opts):
    try:
        outdir = make_output_dir(input_path, parent=OUTPUT_ROOT)
        tag = os.path.basename(outdir)
        _set(job_id, outdir=outdir, tag=tag)

        s1 = os.path.join(outdir, f"step1_segmented_{tag}.json")
        s2 = os.path.join(outdir, f"step2_scores_{tag}.json")
        s3 = os.path.join(outdir, f"step3_candidates_{tag}.json")
        s4 = os.path.join(outdir, f"step4_rules_{tag}.json")
        s5 = os.path.join(outdir, f"step5_final_rules_{tag}.json")
        xlsx = os.path.join(outdir, f"report_{tag}.xlsx")

        step1 = ["step1_segment.py", "--input", input_path, "--output", s1,
                 "--segmenter", opts["segmenter"]]
        if opts["no_phobert_tokenizer"]:
            step1.append("--no-phobert-tokenizer")
        run_cmd(job_id, 0, step1)

        run_cmd(job_id, 1, ["step2_score.py", "--input", s1, "--output", s2])

        run_cmd(job_id, 2, ["step3_candidates.py", "--step1", s1, "--step2", s2,
                            "--output", s3, "--min-gram", str(opts["min_gram"]),
                            "--kw-min-precision", str(opts["kw_min_precision"])])

        step4 = ["step4_select_rules.py", "--step1", s1, "--step3", s3,
                 "--output", s4, "--max-rules", str(opts["max_rules"]),
                 "--max-kw", str(opts["max_kw"]),
                 "--and-cover", str(opts["and_cover"]),
                 "--sim-threshold", str(opts["sim_threshold"])]
        if opts["no_embedding"]:
            step4.append("--no-embedding")
        run_cmd(job_id, 3, step4)

        run_cmd(job_id, 4, ["step5_finalize.py", "--input", s4, "--output", s5,
                            "--max-rules", str(opts["max_rules"]),
                            "--max-kw", str(opts["max_kw"])])

        out6 = run_cmd(job_id, 5, ["step6_export.py", "--step1", s1,
                                   "--step5", s5, "--output", xlsx])

        metrics = parse_metrics(out6)
        rules = load_rules(s5)
        try:                       # bổ sung 'clustering' từ meta step4
            rules["meta"]["clustering"] = read_json(s4)["meta"].get("clustering")
        except Exception:
            pass
        _set(job_id, status="done", report=xlsx, metrics=metrics, rules=rules,
             report_name=os.path.basename(xlsx))
    except Exception as e:
        _set(job_id, status="error", error=str(e))


def parse_metrics(step6_stdout):
    """Bóc precision/recall/f1/acc + TP/FP/TN/FN từ stdout step6."""
    m = {}
    mt = re.search(r"TP=(\d+) FP=(\d+) TN=(\d+) FN=(\d+)", step6_stdout)
    if mt:
        m.update(tp=int(mt[1]), fp=int(mt[2]), tn=int(mt[3]), fn=int(mt[4]))
    mp = re.search(r"Precision=([\d.]+)% Recall=([\d.]+)% F1=([\d.]+)% Acc=([\d.]+)%",
                   step6_stdout)
    if mp:
        m.update(precision=float(mp[1]), recall=float(mp[2]),
                 f1=float(mp[3]), acc=float(mp[4]))
    m["pass_precision"] = m.get("precision", 0) > 99.0
    m["pass_recall"] = m.get("recall", 0) > 30.0
    return m


def load_rules(step5_path):
    from common import kw_display
    data = read_json(step5_path)
    out = {"spam": [], "nonspam": [], "meta": data.get("meta", {})}
    for r in data.get("spam_rules", []):
        out["spam"].append({
            "id": r["id"], "category": kw_display(r.get("category", "")),
            "or": [kw_display(k) for k in r.get("or", [])],
            "and": [kw_display(k) for k in r.get("and", [])],
            "not": [kw_display(k) for k in r.get("not", [])],
        })
    for r in data.get("nonspam_rules", []):
        out["nonspam"].append({
            "id": r["id"], "category": kw_display(r.get("category", "")),
            "or": [kw_display(k) for k in r.get("or", [])],
            "and": [], "not": [],
        })
    return out


# ------------------------------------------------------------------ views ---
@app.route("/")
def index():
    return render_template_string(INDEX_HTML)


@app.route("/upload", methods=["POST"])
def upload():
    mode = request.form.get("mode", "two")

    def _xlsx(f):
        return f and f.filename and f.filename.lower().endswith((".xlsx", ".xls"))

    job_id = uuid.uuid4().hex[:12]
    job_dir = os.path.join(UPLOAD_TMP, job_id)
    os.makedirs(job_dir, exist_ok=True)

    if mode == "one":
        # ---- 1 FILE có sẵn 2 sheet 'spam' & 'nonspam' -> dùng trực tiếp ----
        f_one = request.files.get("file_single")
        if not _xlsx(f_one):
            return "Vui lòng chọn 1 file Excel (.xlsx) có 2 sheet spam & nonspam", 400
        name = (request.form.get("dataset_name") or "").strip() \
            or os.path.splitext(os.path.basename(f_one.filename))[0]
        base = re.sub(r"[^\w\-. ]", "_", name).strip() or "dataset"
        saved = os.path.join(job_dir, f"{base}.xlsx")
        f_one.save(saved)
        try:
            n_spam, n_non = _count_sheets(saved)
        except Exception as e:
            return f"Lỗi đọc file Excel: {e}", 400
        if n_spam == 0:
            return ("File không có sheet tên 'spam' (hoặc sheet spam rỗng). "
                    "Cần 2 sheet: 'spam' & 'nonspam'."), 400
        input_label = (f"{f_one.filename} · 2 sheet "
                       f"({n_spam:,} spam / {n_non:,} nonspam)")
    else:
        # ---- 2 FILE riêng -> gộp thành 1 file 2 sheet ----
        f_spam = request.files.get("file_spam")
        f_non = request.files.get("file_nonspam")
        if not _xlsx(f_spam):
            return "Vui lòng chọn file SPAM (.xlsx)", 400
        if not _xlsx(f_non):
            return "Vui lòng chọn file NON-SPAM (.xlsx)", 400
        name = (request.form.get("dataset_name") or "").strip() \
            or os.path.splitext(os.path.basename(f_spam.filename))[0]
        base = re.sub(r"[^\w\-. ]", "_", name).strip() or "dataset"
        raw_spam = os.path.join(job_dir, "_spam_in.xlsx")
        raw_non = os.path.join(job_dir, "_nonspam_in.xlsx")
        f_spam.save(raw_spam)
        f_non.save(raw_non)
        saved = os.path.join(job_dir, f"{base}.xlsx")
        try:
            n_spam, n_non = build_combined_excel(raw_spam, raw_non, saved)
        except Exception as e:
            return f"Lỗi đọc/gộp file Excel: {e}", 400
        input_label = (f"{f_spam.filename} ({n_spam:,} spam) + "
                       f"{f_non.filename} ({n_non:,} nonspam)")

    def fnum(name, default, cast=float):
        try:
            return cast(request.form.get(name, default))
        except Exception:
            return default

    opts = {
        "segmenter": request.form.get("segmenter", "underthesea"),
        "no_phobert_tokenizer": request.form.get("phobert_tok") != "on",
        "no_embedding": request.form.get("embedding") != "on",
        "min_gram": fnum("min_gram", 3, int),
        "kw_min_precision": fnum("kw_min_precision", 0.97),
        "and_cover": fnum("and_cover", 0.95),
        "sim_threshold": fnum("sim_threshold", 0.62),
        "max_rules": fnum("max_rules", 20, int),
        "max_kw": fnum("max_kw", 20, int),
    }
    with JOBS_LOCK:
        JOBS[job_id] = {
            "status": "running", "input_name": input_label,
            "steps": [{"title": t, "state": "pending", "note": ""} for t in STEP_TITLES],
            "error": None,
        }
    threading.Thread(target=worker, args=(job_id, saved, opts), daemon=True).start()
    return redirect(url_for("result", job_id=job_id))


@app.route("/status/<job_id>")
def status(job_id):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if not job:
            return jsonify({"error": "not found"}), 404
        return jsonify({
            "status": job["status"],
            "steps": job["steps"],
            "error": job.get("error"),
            "metrics": job.get("metrics"),
            "input_name": job.get("input_name"),
            "tag": job.get("tag"),
        })


@app.route("/result/<job_id>")
def result(job_id):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
    if not job:
        abort(404)
    return render_template_string(RESULT_HTML, job_id=job_id, job=job)


@app.route("/download/<job_id>")
def download(job_id):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
    if not job or job.get("status") != "done" or not job.get("report"):
        abort(404)
    return send_file(job["report"], as_attachment=True,
                     download_name=job["report_name"])


# --------------------------------------------------------------- history ---
def _scan_history():
    """Quét OUTPUT_ROOT -> list các lần đã chạy (mỗi folder = 1 lần).

    Mỗi mục: tên folder, thời gian, và các file .xlsx bên trong (report…).
    Sắp xếp mới nhất lên đầu.
    """
    import datetime
    items = []
    try:
        names = os.listdir(OUTPUT_ROOT)
    except OSError:
        return items
    for name in names:
        folder = os.path.join(OUTPUT_ROOT, name)
        if not os.path.isdir(folder):
            continue
        xlsx = [f for f in os.listdir(folder) if f.lower().endswith(".xlsx")]
        if not xlsx:
            continue
        files = []
        for f in sorted(xlsx):
            p = os.path.join(folder, f)
            try:
                files.append({"name": f, "size": os.path.getsize(p)})
            except OSError:
                continue
        try:
            mtime = os.path.getmtime(folder)
        except OSError:
            mtime = 0
        items.append({
            "folder": name,
            "mtime": mtime,
            "when": datetime.datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M")
                    if mtime else "",
            "files": files,
        })
    items.sort(key=lambda x: x["mtime"], reverse=True)
    return items


@app.route("/history")
def history():
    return jsonify({"items": _scan_history()})


@app.route("/history/download/<folder>/<path:filename>")
def history_download(folder, filename):
    # chống path traversal: chỉ cho tên folder/file đơn giản, phải nằm trong OUTPUT_ROOT
    if (os.path.sep in folder or (os.path.altsep and os.path.altsep in folder)
            or folder in ("", ".", "..")):
        abort(404)
    if not filename.lower().endswith((".xlsx", ".xls")):
        abort(404)
    base = os.path.abspath(os.path.join(OUTPUT_ROOT, folder))
    if base != os.path.abspath(OUTPUT_ROOT) and not base.startswith(
            os.path.abspath(OUTPUT_ROOT) + os.path.sep):
        abort(404)
    full = os.path.abspath(os.path.join(base, filename))
    if not (full == base or full.startswith(base + os.path.sep)):
        abort(404)
    if not os.path.isfile(full):
        abort(404)
    return send_file(full, as_attachment=True, download_name=filename)


# ------------------------------------------------------------------ HTML ---
BASE_CSS = """
:root{--bg:#0f172a;--card:#1e293b;--muted:#94a3b8;--fg:#e2e8f0;--accent:#38bdf8;
--ok:#22c55e;--bad:#ef4444;--warn:#f59e0b;--line:#334155}
*{box-sizing:border-box}
body{margin:0;font-family:system-ui,Segoe UI,Roboto,sans-serif;background:var(--bg);
color:var(--fg);line-height:1.5}
.wrap{max-width:1080px;margin:0 auto;padding:28px 20px 60px}
h1{font-size:24px;margin:0 0 4px}.sub{color:var(--muted);margin:0 0 24px;font-size:14px}
a{color:var(--accent)}
.card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:22px;margin-bottom:18px}
label{display:block;font-size:13px;color:var(--muted);margin:10px 0 4px}
input[type=text],input[type=number],select{width:100%;padding:9px 11px;border-radius:9px;
border:1px solid var(--line);background:#0b1220;color:var(--fg);font-size:14px}
.file{border:2px dashed var(--line);border-radius:12px;padding:26px;text-align:center;cursor:pointer;transition:.15s}
.file:hover{border-color:var(--accent)}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px}
.btn{display:inline-block;background:var(--accent);color:#04283a;font-weight:700;border:0;
padding:12px 22px;border-radius:10px;font-size:15px;cursor:pointer;text-decoration:none}
.btn:disabled{opacity:.5;cursor:not-allowed}
.btn.secondary{background:#334155;color:var(--fg)}
details summary{cursor:pointer;color:var(--muted);font-size:13px;margin-top:6px}
.metrics{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px}
.metric{background:#0b1220;border:1px solid var(--line);border-radius:12px;padding:16px}
.metric .v{font-size:26px;font-weight:800}.metric .k{font-size:12px;color:var(--muted)}
.pill{font-size:11px;padding:2px 8px;border-radius:20px;font-weight:700}
.pill.ok{background:rgba(34,197,94,.15);color:var(--ok)}
.pill.bad{background:rgba(239,68,68,.15);color:var(--bad)}
.steps{list-style:none;padding:0;margin:0}
.steps li{display:flex;align-items:center;gap:12px;padding:9px 0;border-bottom:1px solid var(--line);font-size:14px}
.steps li:last-child{border:0}
.dot{width:22px;height:22px;border-radius:50%;flex:0 0 22px;display:grid;place-items:center;font-size:12px;font-weight:700}
.dot.pending{background:#0b1220;color:var(--muted);border:1px solid var(--line)}
.dot.running{background:var(--warn);color:#3a2500}
.dot.done{background:var(--ok);color:#04310f}
.dot.error{background:var(--bad);color:#3a0000}
.note{color:var(--muted);font-size:12px;margin-left:auto}
.rule{background:#0b1220;border:1px solid var(--line);border-radius:12px;padding:14px 16px;margin-bottom:12px}
.rule h4{margin:0 0 8px;font-size:14px}
.rule .rid{color:var(--accent);font-family:ui-monospace,monospace;font-size:12px}
.kw{display:inline-block;background:#1e293b;border:1px solid var(--line);border-radius:7px;
padding:3px 9px;margin:3px 4px 0 0;font-size:12.5px}
.kw.and{border-color:var(--warn);color:#fcd34d}
.kw.not{border-color:var(--bad);color:#fca5a5}
.tag{font-size:11px;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}
.spin{width:16px;height:16px;border:2px solid var(--line);border-top-color:var(--accent);
border-radius:50%;display:inline-block;animation:sp 0.7s linear infinite}
@keyframes sp{to{transform:rotate(360deg)}}
.modes{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:14px}
.modebtn{flex:1;min-width:200px;padding:10px 12px;border-radius:10px;border:1px solid var(--line);
background:#0b1220;color:var(--muted);font-size:13.5px;font-weight:600;cursor:pointer;transition:.15s}
.modebtn:hover{border-color:var(--accent)}
.modebtn.active{background:rgba(56,189,248,.12);border-color:var(--accent);color:var(--fg)}
"""

INDEX_HTML = """<!doctype html><html lang=vi><head><meta charset=utf-8>
<meta name=viewport content="width=device-width,initial-scale=1">
<title>Rulebase Spam · Import Excel</title><style>""" + BASE_CSS + """</style></head>
<body><div class=wrap>
<h1>🛡️ Hệ thống sinh Rulebase chống spam</h1>
<p class=sub>Import dữ liệu Excel (dòng 1 là tiêu đề, nội dung từ <b>dòng 2</b>, lấy
cột <code>content</code>/<code>nội dung</code> hoặc cột đầu) → tự sinh rulebase
(OR/AND/NOT) đảm bảo <b>precision&gt;99%</b> → tải report về.</p>
<form class=card method=post action="/upload" enctype="multipart/form-data" id=frm>
  <input type=hidden name=mode id=mode value=two>
  <div class=modes>
    <button type=button class="modebtn" data-mode=two id=mb_two>📑 2 file riêng (spam + nonspam)</button>
    <button type=button class="modebtn" data-mode=one id=mb_one>📘 1 file (2 sheet: spam &amp; nonspam)</button>
  </div>

  <div id=pane_two>
  <div class=grid style="grid-template-columns:repeat(auto-fit,minmax(240px,1fr))">
    <div>
      <label>① File SPAM (.xlsx)</label>
      <label class=file id=drop_spam>
        <div id=fname_spam>🚫 Bấm để chọn file SPAM …</div>
        <input type=file name=file_spam accept=".xlsx,.xls" hidden id=file_spam>
      </label>
    </div>
    <div>
      <label>② File NON-SPAM (.xlsx)</label>
      <label class=file id=drop_non>
        <div id=fname_non>✅ Bấm để chọn file NON-SPAM …</div>
        <input type=file name=file_nonspam accept=".xlsx,.xls" hidden id=file_nonspam>
      </label>
    </div>
  </div>
  </div>

  <div id=pane_one style=display:none>
    <label>File Excel có <b>2 sheet</b>: <code>spam</code> &amp; <code>nonspam</code></label>
    <label class=file id=drop_single>
      <div id=fname_single>📘 Bấm để chọn file .xlsx (2 sheet) …</div>
      <input type=file name=file_single accept=".xlsx,.xls" hidden id=file_single>
    </label>
  </div>

  <label style=margin-top:12px>Tên dataset (đặt tên folder kết quả — có thể bỏ trống)</label>
  <input type=text name=dataset_name placeholder="vd: nguyenkim">
  <details style=margin-top:8px><summary>⚙️ Tuỳ chọn nâng cao (có thể bỏ qua)</summary>
  <div class=grid style=margin-top:10px>
    <div><label>Bộ tách từ</label><select name=segmenter>
      <option value=underthesea selected>underthesea (nhanh)</option>
      <option value=auto>auto (VnCoreNLP nếu có Java)</option>
      <option value=whitespace>whitespace (thô)</option></select></div>
    <div><label>min_gram (số từ/keyword)</label><input type=number name=min_gram value=3 min=1 max=6></div>
    <div><label>kw_min_precision</label><input type=number step=0.01 name=kw_min_precision value=0.97 min=0.5 max=1></div>
    <div><label>and_cover (AND phủ ≥)</label><input type=number step=0.05 name=and_cover value=0.95 min=0.5 max=1></div>
    <div><label>sim_threshold (gom nghĩa)</label><input type=number step=0.02 name=sim_threshold value=0.62 min=0.3 max=0.9></div>
    <div><label>max_rules</label><input type=number name=max_rules value=20 min=1 max=100></div>
    <div><label>max_kw / rule</label><input type=number name=max_kw value=20 min=1 max=100></div>
  </div>
  <label style=margin-top:12px><input type=checkbox name=embedding checked> Gom nhóm bằng PhoBERT embedding (ngữ nghĩa)</label>
  <label><input type=checkbox name=phobert_tok> Nạp PhoBERT tokenizer ở B1 (chậm hơn, không bắt buộc)</label>
  </details>
  <div style=margin-top:18px><button class=btn type=submit id=go>🚀 Chạy pipeline</button></div>
</form>
<p class=sub>Kết quả mỗi lần chạy lưu vào 1 folder riêng (tên = tên dataset, tự thêm _v1/_v2 nếu trùng); trên Docker là <code>./output/</code> ở máy host.</p>

<div class=card>
  <div style="display:flex;align-items:center;gap:10px">
    <div class=tag>📚 Danh mục đã chạy (report Excel đã lưu)</div>
    <button type=button class="btn secondary" id=refhist
      style="margin-left:auto;padding:6px 12px;font-size:12px">⟳ Làm mới</button>
  </div>
  <div id=history style=margin-top:12px><span class=note>Đang tải danh mục…</span></div>
</div>
</div>
<script>
async function loadHistory(){
  const box=document.getElementById('history');
  try{
    const r=await fetch('/history');const d=await r.json();
    const items=d.items||[];
    if(!items.length){box.innerHTML='<span class=note>Chưa có lần chạy nào được lưu.</span>';return;}
    box.innerHTML=items.map(it=>{
      const files=it.files.map(f=>{
        const kb=(f.size/1024).toFixed(0);
        const href='/history/download/'+encodeURIComponent(it.folder)+'/'+encodeURIComponent(f.name);
        return `<a class=kw href="${href}">⬇️ ${esc(f.name)} <span class=note style="margin:0">(${kb} KB)</span></a>`;
      }).join('');
      return `<div class=rule><h4><span class=rid>📁 ${esc(it.folder)}</span>
        <span class=note style="margin-left:8px">${esc(it.when)}</span></h4>${files}</div>`;
    }).join('');
  }catch(e){box.innerHTML='<span class=note>Không tải được danh mục.</span>';}
}
function esc(s){return (s||'').replace(/</g,'&lt;').replace(/>/g,'&gt;')}
document.getElementById('refhist').onclick=loadHistory;
loadHistory();
function wire(inputId,labelId,icon){
  const fi=document.getElementById(inputId),fn=document.getElementById(labelId);
  fi.onchange=()=>{if(fi.files[0])fn.textContent=icon+' '+fi.files[0].name};
}
wire('file_spam','fname_spam','🚫');
wire('file_nonspam','fname_non','✅');
wire('file_single','fname_single','📘');

const modeInput=document.getElementById('mode');
function setMode(m){
  modeInput.value=m;
  document.getElementById('pane_two').style.display = m=='two'?'':'none';
  document.getElementById('pane_one').style.display = m=='one'?'':'none';
  document.getElementById('mb_two').classList.toggle('active',m=='two');
  document.getElementById('mb_one').classList.toggle('active',m=='one');
}
document.getElementById('mb_two').onclick=()=>setMode('two');
document.getElementById('mb_one').onclick=()=>setMode('one');
setMode('two');

document.getElementById('frm').onsubmit=(e)=>{
  const m=modeInput.value;
  if(m=='two'){
    if(!document.getElementById('file_spam').files[0]||!document.getElementById('file_nonspam').files[0]){
      e.preventDefault();alert('Chọn đủ 2 file: SPAM và NON-SPAM.');return;}
  }else{
    if(!document.getElementById('file_single').files[0]){
      e.preventDefault();alert('Chọn 1 file Excel (2 sheet spam & nonspam).');return;}
  }
  document.getElementById('go').disabled=true;
  document.getElementById('go').innerHTML='<span class=spin></span> Đang tải lên…';
};
</script></body></html>"""

RESULT_HTML = """<!doctype html><html lang=vi><head><meta charset=utf-8>
<meta name=viewport content="width=device-width,initial-scale=1">
<title>Kết quả · Rulebase</title><style>""" + BASE_CSS + """</style></head>
<body><div class=wrap>
<h1>📊 Kết quả xử lý</h1>
<p class=sub>File: <b>{{ job.input_name }}</b> · <a href="/">⟵ chạy file khác</a></p>
<div class=card>
  <div class=tag>Tiến độ</div>
  <ul class=steps id=steps></ul>
</div>
<div id=result></div>
</div>
<script>
const jid="{{ job_id }}";
const ICON={pending:'',running:'●',done:'✓',error:'✕'};
function render(d){
  const ul=document.getElementById('steps');ul.innerHTML='';
  d.steps.forEach(s=>{
    const li=document.createElement('li');
    li.innerHTML=`<span class="dot ${s.state}">${s.state=='running'?'<span class=spin></span>':ICON[s.state]}</span>
      <span>${s.title}</span><span class=note>${s.note||''}</span>`;
    ul.appendChild(li);
  });
  const box=document.getElementById('result');
  if(d.status=='error'){
    box.innerHTML=`<div class=card style="border-color:var(--bad)"><b style=color:var(--bad)>Lỗi:</b>
      <pre style="white-space:pre-wrap;color:#fca5a5">${(d.error||'').replace(/</g,'&lt;')}</pre></div>`;
    return true;
  }
  if(d.status=='done'){
    const m=d.metrics||{};
    box.innerHTML=`
    <div class=card><div class=tag>Chỉ số trên tập test (25% giữ riêng)</div>
      <div class=metrics style=margin-top:12px>
        <div class=metric><div class=k>Precision</div><div class=v>${(m.precision||0).toFixed(2)}%</div>
          <span class="pill ${m.pass_precision?'ok':'bad'}">${m.pass_precision?'PASS >99%':'FAIL'}</span></div>
        <div class=metric><div class=k>Recall</div><div class=v>${(m.recall||0).toFixed(2)}%</div>
          <span class="pill ${m.pass_recall?'ok':'bad'}">${m.pass_recall?'PASS >30%':'FAIL'}</span></div>
        <div class=metric><div class=k>F1-score</div><div class=v>${(m.f1||0).toFixed(2)}%</div></div>
        <div class=metric><div class=k>TP / FP</div><div class=v>${m.tp||0} / ${m.fp||0}</div>
          <div class=k>FN=${m.fn||0} · TN=${m.tn||0}</div></div>
      </div>
      <div style=margin-top:18px>
        <a class=btn href="/download/${jid}">⬇️ Tải report Excel (${d.tag})</a>
      </div>
    </div>
    <div class=card><div class=tag>Rulebase — bấm để tải bản đầy đủ trong Excel</div>
      <div id=rules style=margin-top:12px><span class=note>Đang tải rulebase…</span></div>
    </div>`;
    loadRules();
    return true;
  }
  return false;
}
async function loadRules(){
  const r=await fetch('/rules/'+jid);if(!r.ok)return;const rb=await r.json();
  const el=document.getElementById('rules');el.innerHTML='';
  el.insertAdjacentHTML('beforeend',`<p class=note>Gom nhóm: <b>${rb.meta.clustering||'-'}</b> ·
    ${rb.spam.length} rule spam · ${rb.nonspam.length} rule nonspam</p>`);
  rb.spam.forEach(ru=>{
    const orh=ru.or.map(k=>`<span class=kw>${esc(k)}</span>`).join('');
    const andh=ru.and.map(k=>`<span class="kw and">AND ${esc(k)}</span>`).join('');
    const noth=ru.not.map(k=>`<span class="kw not">NOT ${esc(k)}</span>`).join('');
    el.insertAdjacentHTML('beforeend',
      `<div class=rule><h4><span class=rid>${ru.id}</span> · ${esc(ru.category)}</h4>
       ${orh}${andh}${noth}</div>`);
  });
}
function esc(s){return (s||'').replace(/</g,'&lt;').replace(/>/g,'&gt;')}
async function poll(){
  try{const r=await fetch('/status/'+jid);const d=await r.json();
    if(render(d))return;}catch(e){}
  setTimeout(poll,1200);
}
poll();
</script></body></html>"""


@app.route("/rules/<job_id>")
def rules(job_id):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
    if not job or job.get("status") != "done":
        return jsonify({"error": "not ready"}), 404
    return jsonify(job.get("rules", {"spam": [], "nonspam": [], "meta": {}}))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--host", default="127.0.0.1")
    ap.add_argument("--port", type=int, default=5000)
    ap.add_argument("--debug", action="store_true")
    args = ap.parse_args()
    print(f"\n  Rulebase web đang chạy tại:  http://{args.host}:{args.port}\n")
    app.run(host=args.host, port=args.port, debug=args.debug, threaded=True)


if __name__ == "__main__":
    main()
